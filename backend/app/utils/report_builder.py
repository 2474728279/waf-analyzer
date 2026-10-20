"""专业Word报告生成器 - WAF日志分析报告"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from datetime import datetime
from typing import Optional


class WAFReportBuilder:
    def __init__(self):
        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self):
        """设置文档样式。"""
        style = self.doc.styles["Normal"]
        font = style.font
        font.name = "Microsoft YaHei"
        font.size = Pt(10.5)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

        for level in range(1, 4):
            heading_style = self.doc.styles[f"Heading {level}"]
            heading_font = heading_style.font
            heading_font.name = "Microsoft YaHei"
            heading_style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            heading_font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    def build_report(self, project: dict, analysis: dict) -> Document:
        """构建完整的WAF日志分析报告。"""
        self._add_cover_page(project, analysis)
        self._add_overview(project, analysis)
        self._add_attack_distribution(analysis)
        self._add_root_cause_analysis(analysis)
        self._add_proposed_rules(analysis)
        self._add_deployment_guide(analysis)
        return self.doc

    def _add_cover_page(self, project: dict, analysis: dict):
        """添加封面页。"""
        for _ in range(6):
            self.doc.add_paragraph("")

        title = self.doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("WAF日志分析报告")
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        self.doc.add_paragraph("")

        subtitle = self.doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        customer = project.get("customer_name", "")
        if customer:
            run = subtitle.add_run(f"{customer}")
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x4A, 0x4A, 0x6A)

        self.doc.add_paragraph("")
        self.doc.add_paragraph("")

        info = self.doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        device_name = project.get("waf_device_name", "")
        device_ip = project.get("waf_device_ip", "")
        run = info.add_run(f"WAF设备：{device_name}（{device_ip}）")
        run.font.size = Pt(12)

        date_para = self.doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = date_para.add_run(f"报告日期：{datetime.now().strftime('%Y年%m月%d日')}")
        run.font.size = Pt(12)

        self.doc.add_page_break()

    def _add_overview(self, project: dict, analysis: dict):
        """添加分析概述。"""
        self.doc.add_heading("1. 分析概述", level=1)

        total = analysis.get("total_logs", 0)
        scanner_count = analysis.get("scanner_log_count", 0)
        non_scanner = total - scanner_count
        rules = analysis.get("proposed_rules", [])
        time_start = analysis.get("time_range_start", "")
        time_end = analysis.get("time_range_end", "")

        table = self.doc.add_table(rows=6, cols=2, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        data = [
            ("客户名称", project.get("customer_name", "")),
            ("WAF设备", f"{project.get('waf_device_name', '')}（{project.get('waf_device_ip', '')}）"),
            ("分析时间范围", f"{time_start} ~ {time_end}" if time_start else "全部"),
            ("日志总量", f"{total:,}"),
            ("扫描器相关告警", f"{scanner_count:,}（{scanner_count/total*100:.1f}%）" if total > 0 else "0"),
            ("建议规则数", f"{len(rules)}"),
        ]

        for i, (key, val) in enumerate(data):
            table.rows[i].cells[0].text = key
            table.rows[i].cells[1].text = str(val)
            for cell in table.rows[i].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10.5)

        self.doc.add_paragraph("")

    def _add_attack_distribution(self, analysis: dict):
        """添加攻击类型分布。"""
        self.doc.add_heading("2. 攻击类型分布", level=1)

        dist = analysis.get("attack_type_distribution", {})
        total = analysis.get("total_logs", 0)

        if not dist:
            self.doc.add_paragraph("无攻击类型分布数据。")
            return

        table = self.doc.add_table(rows=len(dist) + 1, cols=3, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["攻击类型", "告警数", "占比"]
        for j, h in enumerate(headers):
            cell = table.rows[0].cells[j]
            cell.text = h
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.bold = True

        for i, (attack_type, count) in enumerate(sorted(dist.items(), key=lambda x: -x[1]), 1):
            pct = f"{count / total * 100:.1f}%" if total > 0 else "0%"
            table.rows[i].cells[0].text = attack_type
            table.rows[i].cells[1].text = f"{count:,}"
            table.rows[i].cells[2].text = pct

        self.doc.add_paragraph("")

    def _add_root_cause_analysis(self, analysis: dict):
        """添加根因分析。"""
        self.doc.add_heading("3. 根因分析", level=1)

        root_causes = analysis.get("root_causes", [])
        if not root_causes:
            self.doc.add_paragraph("无根因分析数据。")
            return

        table = self.doc.add_table(rows=len(root_causes) + 1, cols=4, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["根因类别", "告警数", "占比", "描述"]
        for j, h in enumerate(headers):
            cell = table.rows[0].cells[j]
            cell.text = h
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.bold = True

        for i, rc in enumerate(root_causes, 1):
            table.rows[i].cells[0].text = rc.get("category", "")
            table.rows[i].cells[1].text = str(rc.get("log_count", 0))
            pct = rc.get("percentage", 0)
            table.rows[i].cells[2].text = f"{pct:.1f}%" if isinstance(pct, (int, float)) else str(pct)
            table.rows[i].cells[3].text = rc.get("description", "")

            # 建议操作
            if rc.get("suggested_action"):
                desc = table.rows[i].cells[3].text
                table.rows[i].cells[3].text = f"{desc}\n建议：{rc['suggested_action']}"

        self.doc.add_paragraph("")

    def _add_proposed_rules(self, analysis: dict):
        """添加建议规则。"""
        self.doc.add_heading("4. 建议规则", level=1)

        rules = analysis.get("proposed_rules", [])
        if not rules:
            self.doc.add_paragraph("无建议规则。")
            return

        self.doc.add_paragraph(f"共生成 {len(rules)} 条规则，覆盖告警情况如下：")

        # 规则汇总表
        table = self.doc.add_table(rows=min(len(rules), 50) + 1, cols=5, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["序号", "规则名称", "关闭模块", "覆盖告警数", "策略"]
        for j, h in enumerate(headers):
            cell = table.rows[0].cells[j]
            cell.text = h
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.bold = True

        for i, rule in enumerate(rules[:50], 1):
            modules = ", ".join(rule.get("disabled_modules", []))
            table.rows[i].cells[0].text = str(i)
            table.rows[i].cells[1].text = rule.get("name", "")
            table.rows[i].cells[2].text = modules
            table.rows[i].cells[3].text = str(rule.get("coverage_count", 0))
            table.rows[i].cells[4].text = rule.get("strategy", "")

        self.doc.add_paragraph("")

        # 规则详情
        self.doc.add_heading("4.1 规则详情", level=2)
        for i, rule in enumerate(rules[:30], 1):
            self.doc.add_heading(f"规则 {i}: {rule.get('name', '')}", level=3)

            detail_table = self.doc.add_table(rows=6, cols=2, style="Light Grid Accent 1")
            details = [
                ("规则描述", rule.get("comment", "")),
                ("攻击类型", rule.get("attack_type", "")),
                ("关闭检测模块", ", ".join(rule.get("disabled_modules", []))),
                ("匹配策略", rule.get("strategy", "")),
                ("覆盖告警数", str(rule.get("coverage_count", 0))),
                ("置信度", f"{rule.get('confidence', 0):.0%}"),
            ]
            for j, (key, val) in enumerate(details):
                detail_table.rows[j].cells[0].text = key
                detail_table.rows[j].cells[1].text = val

            self.doc.add_paragraph("")

    def _add_deployment_guide(self, analysis: dict):
        """添加部署指南。"""
        self.doc.add_heading("5. 部署指南", level=1)

        rules = analysis.get("proposed_rules", [])
        enabled_count = sum(1 for r in rules if r.get("is_enabled"))

        self.doc.add_paragraph(
            "所有建议规则默认为禁用状态，需经人工审阅确认后方可启用部署。"
        )

        guide_items = [
            "1. 审阅规则：逐条检查规则名称、描述、匹配条件和关闭模块是否正确",
            "2. 修改规则：对不满足需求的规则进行调整（如修改URL匹配范围、增减关闭模块）",
            "3. 启用规则：确认无误后，在WAF管理界面或通过API启用规则",
            "4. 观察验证：启用后持续观察WAF告警变化，确认误报消除效果",
            "5. 持续调优：发现新的误报模式时，创建新规则或修改现有规则",
        ]

        for item in guide_items:
            para = self.doc.add_paragraph(item)
            para.paragraph_format.space_after = Pt(4)

    def save(self, path: str):
        """保存报告到文件。"""
        self.doc.save(path)


def generate_word_report(project: dict, analysis: dict, output_path: str) -> str:
    """生成Word格式WAF日志分析报告。"""
    builder = WAFReportBuilder()
    builder.build_report(project, analysis)
    builder.save(output_path)
    return output_path
